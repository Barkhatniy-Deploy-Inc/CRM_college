import os
from docxtpl import DocxTemplate
from io import BytesIO
from datetime import datetime

# Каталог сервиса techcard (backend/techcard) и шаблон документа.
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_TEMPLATE_PATH = os.path.join(
    BASE_DIR, "templates", "Tekhnologicheskaia-karta-zaniatiia.docx"
)

def generate_techcard_docx(card_data: dict) -> BytesIO:
    """
    Генерирует .docx файл на основе данных техкарты.
    card_data - это словарь с данными из БД и вложенными этапами.
    """
    # Путь к шаблону (можно переопределить через TECHCARD_TEMPLATE_PATH)
    template_path = os.getenv("TECHCARD_TEMPLATE_PATH", DEFAULT_TEMPLATE_PATH)
    
    # Если шаблона нет, мы могли бы создать его программно, 
    # но docxtpl требует файл. Для начала создадим базовый контекст.
    
    context = {
        'tema': card_data.get('tema', 'Не указана'),
        'nomer': card_data.get('nomer_zanyatiya', '1'),
        'cel': card_data.get('cel_zanyatiya', ''),
        'obuch': card_data.get('zadachi_obuch', ''),
        'razv': card_data.get('zadachi_razv', ''),
        'vosp': card_data.get('zadachi_vosp', ''),
        'tech': card_data.get('ped_tech', ''),
        'res': card_data.get('prognoz_result', ''),
        'oborud': card_data.get('oborudovanie', ''),
        'istoch': card_data.get('istochniki', ''),
        'stages': card_data.get('stages', []),
        'date': datetime.now().strftime("%d.%m.%Y")
    }

    # Создаем временный файл в памяти
    target_stream = BytesIO()
    
    try:
        if os.path.exists(template_path):
            doc = DocxTemplate(template_path)
            doc.render(context)
            doc.save(target_stream)
        else:
            # Если шаблона нет, создаем простейший документ через python-docx
            from docx import Document
            doc = Document()
            doc.add_heading(f"Технологическая карта занятия: {context['tema']}", 0)
            doc.add_paragraph(f"Цель: {context['cel']}")
            doc.add_heading("Ход занятия", 1)
            for s in context['stages']:
                doc.add_paragraph(f"Этап: {s.get('nazvanie_etapa')}\nПрод: {s.get('dlitelnost')} мин.\nДеят. преп.: {s.get('deyatelnost_prepod')}")
            doc.save(target_stream)
            
        target_stream.seek(0)
        return target_stream
    except Exception as e:
        print(f"Error generating DOCX: {e}")
        raise e
