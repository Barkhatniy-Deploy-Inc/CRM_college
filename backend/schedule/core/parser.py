import openpyxl
from typing import List, Dict, Optional, Tuple
import re
from datetime import datetime, time
import logging

logger = logging.getLogger(__name__)

# Сопоставление дней недели
DAY_NAMES_RUSSIAN = {
    'ПОНЕДЕЛЬНИК': 1, 'пн': 1,
    'ВТОРНИК': 2, 'вт': 2,
    'СРЕДА': 3, 'ср': 3,
    'ЧЕТВЕРГ': 4, 'чт': 4,
    'ПЯТНИЦА': 5, 'пт': 5,
    'СУББОТА': 6, 'сб': 6,
    'ВОСКРЕСЕНЬЕ': 7, 'вс': 7
}


def parse_excel_schedule(file_path: str) -> List[Dict]:
    """
    Парсинг Excel файла с расписанием занятий СурГУ.
    
    Структура файла:
    - Строки 1-5: заголовок документа
    - Строка 6: шапка с названиями групп и колонок
      Колонки: [0]=пусто | [1]=Дата | [2]=Занятие | [3]=Время | [4]=Группа1 | [5]=Ауд1 | [6]=Группа2 | [7]=Ауд2 | ...
    - Остальные строки: дни недели с занятиями
    
    Возвращает список занятий с полной информацией.
    """
    try:
        workbook = openpyxl.load_workbook(file_path, data_only=True)
    except Exception as e:
        logger.error(f"Ошибка при открытии файла {file_path}: {e}")
        raise

    schedule_entries = []
    
    # Обрабатываем каждый лист в файле (каждый лист = один курс)
    for sheet_name in workbook.sheetnames:
        if sheet_name.upper() == 'ЗВОНКИ':  # Пропускаем служебные листы
            continue
            
        try:
            sheet = workbook[sheet_name]
            logger.info(f"Обработка листа: {sheet_name}")
            
            # Извлекаем даты из заголовка (строка 5)
            date_info = _extract_date_range(sheet)
            
            # Извлекаем названия групп из шапки таблицы (строка 6)
            groups_info = _extract_groups_header(sheet)
            
            if not groups_info:
                logger.warning(f"Не найдены группы в листе {sheet_name}")
                continue
            
            # Обрабатываем строки с занятиями (начиная со строки 7)
            sheet_entries = _parse_sheet_content(sheet, groups_info, date_info, sheet_name)
            schedule_entries.extend(sheet_entries)
            
        except Exception as e:
            logger.error(f"Ошибка при обработке листа {sheet_name}: {e}", exc_info=True)
            continue
    
    logger.info(f"Успешно спарсено {len(schedule_entries)} занятий")
    return schedule_entries


def _extract_date_range(sheet) -> Optional[Dict[str, str]]:
    """Извлекает диапазон дат из заголовка документа (строка 5)"""
    try:
        header_row = list(sheet.iter_rows(min_row=5, max_row=5, values_only=True))[0]
        if header_row and len(header_row) > 1 and header_row[1]:
            date_text = str(header_row[1]).strip()
            # Формат: "с 01 сентября 2025 по 06 сентября 2025"
            match = re.search(r'(\d{2}\s+\w+\s+\d{4})\s+по\s+(\d{2}\s+\w+\s+\d{4})', date_text)
            if match:
                return {
                    'start_date': match.group(1),
                    'end_date': match.group(2),
                    'original': date_text
                }
    except Exception as e:
        logger.debug(f"Не удалось извлечь даты: {e}")
    
    return None


def _extract_groups_header(sheet) -> List[Dict[str, str]]:
    """
    Извлекает названия групп из шапки таблицы (строка 6).
    
    Формат: [0]=пусто | [1]=Дата | [2]=Занятие | [3]=Время | [4]=Группа1 | [5]=Ауд1 | [6]=Группа2 | [7]=Ауд2 | ...
    Возвращает список: [{'name': 'ПД-25/9-П', 'col_group': 4, 'col_room': 5}, ...]
    """
    groups_info = []
    try:
        header_row = list(sheet.iter_rows(min_row=6, max_row=6, values_only=True))[0]
        
        # Пропускаем первые 4 колонки: [0]=пусто, [1]=Дата, [2]=Занятие, [3]=Время
        # Названия групп начинаются с колонки 4
        i = 4
        while i < len(header_row):
            group_name = header_row[i] if i < len(header_row) else None
            
            # Очищаем название группы от пробелов
            if group_name and isinstance(group_name, str):
                group_name = group_name.strip()
                
                # Проверяем, что это название группы (содержит буквы и цифры)
                # Примеры: ПД-25/9-П, СА-25/9-П, ЭРОЭ-25/9-П, СД-25/9-П
                if re.search(r'[А-ЯЁ].*\d', group_name) and 'ауд' not in group_name.lower():
                    groups_info.append({
                        'name': group_name,
                        'col_group': i,
                        'col_room': i + 1 if i + 1 < len(header_row) else None
                    })
                    i += 2  # Группа + аудитория
                else:
                    i += 1
            else:
                i += 1
                
    except Exception as e:
        logger.debug(f"Ошибка при извлечении заголовков групп: {e}")
    
    return groups_info


def _parse_sheet_content(sheet, groups_info: List[Dict], date_info: Optional[Dict], 
                         sheet_name: str) -> List[Dict]:
    """Парсит содержимое листа и возвращает список занятий"""
    entries = []
    current_date_str = None
    current_day_of_week = None
    
    # Начинаем со строки 7 (первые 6 строк - заголовок)
    for row_idx, row in enumerate(sheet.iter_rows(min_row=7, values_only=True), start=7):
        if not row or not any(row):
            continue
        
        # Проверяем день недели в колонке 1
        col1_text = str(row[1] or '').strip() if len(row) > 1 and row[1] else ''
        
        # Ищем день недели в колонке 1
        day_match = _extract_day_of_week(col1_text)
        if day_match:
            current_day_of_week = day_match['day_num']
            current_date_str = day_match['date_str']
            continue
        
        # Обрабатываем строку с занятиями (содержит время в колонке 3)
        if current_date_str and current_day_of_week:
            time_slot = _extract_time_slot(row)
            pair_num = _extract_pair_number(row)
            
            if time_slot:
                # Для каждой группы извлекаем информацию о занятии
                for group_info in groups_info:
                    col_group = group_info['col_group']
                    col_room = group_info['col_room']
                    
                    if col_group < len(row):
                        subject_and_teacher = str(row[col_group] or '').strip()
                        room = str(row[col_room] or '').strip() if col_room and col_room < len(row) else None
                        
                        if subject_and_teacher and subject_and_teacher != '-':
                            # Парсим предмет и преподавателя
                            subject, teacher = _parse_subject_and_teacher(subject_and_teacher)
                            
                            entry = {
                                'group_name': group_info['name'],
                                'date': current_date_str,
                                'day_of_week': current_day_of_week,
                                'time_slot': time_slot,
                                'pair_number': pair_num,
                                'subject': subject,
                                'teacher': teacher,
                                'auditorium': room if room else None,
                                'sheet_name': sheet_name
                            }
                            entries.append(entry)
    
    return entries


def _extract_day_of_week(text: str) -> Optional[Dict]:
    """
    Извлекает день недели и дату из текста.
    Формат: "ПОНЕДЕЛЬНИК                                 01.09.2025"
    """
    for day_name, day_num in DAY_NAMES_RUSSIAN.items():
        if text.startswith(day_name):
            # Ищем дату в формате ДД.MM.ГГГГ
            date_match = re.search(r'(\d{2}\.\d{2}\.\d{4})', text)
            if date_match:
                return {
                    'day_name': day_name,
                    'day_num': day_num,
                    'date_str': date_match.group(1)
                }
    
    return None


def _extract_time_slot(row) -> Optional[str]:
    """
    Извлекает временной слот из колонки 3 (индекс 3).
    Формат: "09:40-10:40"
    """
    if len(row) > 3 and row[3]:
        time_str = str(row[3]).strip()
        # Проверяем формат ЧЧ:ММ-ЧЧ:ММ
        if re.match(r'\d{2}:\d{2}-\d{2}:\d{2}', time_str):
            return time_str
    
    return None


def _extract_pair_number(row) -> Optional[int]:
    """Извлекает номер пары из колонки 2"""
    if len(row) > 2 and row[2]:
        try:
            pair_num = int(row[2])
            if 0 < pair_num <= 8:
                return pair_num
        except (ValueError, TypeError):
            pass
    
    return None


def _parse_subject_and_teacher(text: str) -> Tuple[str, Optional[str]]:
    """
    Парсит строку, содержащую предмет и преподавателя.
    Формат может быть:
    - "Математика Туманова И.С."
    - "Математика                                        Туманова И.С."
    - "Иностранный язык Аллаярова Ю.Ф."
    
    Возвращает кортеж (предмет, преподаватель)
    """
    # Очищаем лишние пробелы
    text = re.sub(r'\s+', ' ', text).strip()
    
    # Пытаемся найти фамилию преподавателя в конце
    # Формат: "Фамилия И.О." или "Фамилия И.В."
    teacher_match = re.search(r'\s([А-ЯЁ][а-яё]+\s+[А-ЯЁ]\.[А-ЯЁ]\.)', text)
    
    if teacher_match:
        teacher = teacher_match.group(1).strip()
        subject = text[:teacher_match.start()].strip()
        return subject, teacher
    
    # Если не нашли преподавателя, возвращаем весь текст как предмет
    return text, None
