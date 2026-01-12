#!/usr/bin/env python3
"""Отладка парсера"""

import openpyxl
from pathlib import Path

schedule_path = Path('РАСПИСАНИЕ')
xlsx_files = sorted([f for f in schedule_path.glob('*.xlsx') if not f.name.startswith('~')])

file_path = str(xlsx_files[0])
print(f'Анализирую: {xlsx_files[0].name}')

wb = openpyxl.load_workbook(file_path, data_only=True)
sheet = wb[wb.sheetnames[0]]

print(f'Лист: {sheet.title}')
print()

# Проверяем строку 6 (шапка)
print('Строка 6 (шапка с группами):')
row6 = list(sheet.iter_rows(min_row=6, max_row=6, values_only=True))[0]
for i, cell in enumerate(row6):
    if cell:
        print(f'  Col {i}: {cell}')

print()
print('Строки 7-15 (примеры данных):')
for idx, row in enumerate(sheet.iter_rows(min_row=7, max_row=15, values_only=True), 7):
    print(f'Row {idx}: {row}')

print()
print('Проверяем логику извлечения дня недели и времени:')
# Проверяем строку 7
row7 = list(sheet.iter_rows(min_row=7, max_row=7, values_only=True))[0]
print(f'Row 7[0]: {row7[0]}')
print(f'Row 7[2]: {row7[2]}')
print(f'Row 7[3]: {row7[3]}')

# Проверяем содержимое ячеек
print()
print('Проверяем, какие данные в ячейках для групп:')
for col in range(3, 11):
    print(f'Col {col}: {row7[col]}')
