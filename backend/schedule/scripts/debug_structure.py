#!/usr/bin/env python3
"""Детальная отладка структуры"""

import openpyxl
from pathlib import Path
import re

schedule_path = Path('РАСПИСАНИЕ')
xlsx_files = sorted([f for f in schedule_path.glob('*.xlsx') if not f.name.startswith('~')])

file_path = str(xlsx_files[0])
print(f'Анализирую: {xlsx_files[0].name}\n')

wb = openpyxl.load_workbook(file_path, data_only=True)
sheet = wb[wb.sheetnames[0]]

print('=' * 120)
print('ПРОВЕРКА СТРУКТУРЫ')
print('=' * 120)

# Проверяем, какие данные в какой колонке
for row_idx in range(6, 20):
    row = list(sheet.iter_rows(min_row=row_idx, max_row=row_idx, values_only=True))[0]
    print(f'\nСтрока {row_idx}:')
    
    # Берем только первые 12 колонок
    for col_idx in range(min(12, len(row))):
        val = row[col_idx]
        if val is not None:
            val_str = str(val)
            if len(val_str) > 50:
                val_str = val_str[:47] + '...'
            print(f'  Col {col_idx:2d}: {val_str}')
        else:
            print(f'  Col {col_idx:2d}: None')
    
    # Проверяем день недели в col 1
    if row[1]:
        first_cell = str(row[1]).strip()
        day_names = ['ПОНЕДЕЛЬНИК', 'ВТОРНИК', 'СРЕДА', 'ЧЕТВЕРГ', 'ПЯТНИЦА', 'СУББОТА', 'ВОСКРЕСЕНЬЕ']
        for day in day_names:
            if first_cell.startswith(day):
                print(f'  >>> НАЙДЕН ДЕНЬ НЕДЕЛИ: {day}')
                date_match = re.search(r'(\d{2}\.\d{2}\.\d{4})', first_cell)
                if date_match:
                    print(f'  >>> НАЙДЕНА ДАТА: {date_match.group(1)}')
